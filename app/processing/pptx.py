"""
Spec section 30: extract slide title/text, notes if available and
intentionally supported (they are, here), include slide number.
"""
from __future__ import annotations

import io

from pptx import Presentation

from app.processing.base import ExtractedContent
from app.sharepoint.files import DownloadedFile


class PptxProcessor:
    async def extract(self, file: DownloadedFile) -> ExtractedContent:
        try:
            presentation = Presentation(io.BytesIO(file.content))
        except Exception as exc:
            return ExtractedContent(text="", warnings=[f"failed to open .pptx: {exc}"])

        parts: list[str] = []
        slide_count = 0

        for idx, slide in enumerate(presentation.slides, start=1):
            slide_count = idx
            parts.append(f"\n--- SLIDE {idx} ---")

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        parts.append(text)

            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                if notes_frame is not None:
                    notes_text = notes_frame.text.strip()
                    if notes_text:
                        parts.append(f"[NOTES] {notes_text}")

        return ExtractedContent(text="\n".join(parts), metadata={"slide_count": slide_count})
